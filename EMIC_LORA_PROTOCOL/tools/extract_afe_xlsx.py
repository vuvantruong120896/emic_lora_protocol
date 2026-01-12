from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import openpyxl


@dataclass(frozen=True)
class SheetPreview:
    name: str
    max_row: int
    max_col: int
    non_empty_rows: list[tuple[int, list[str]]]


def _cell_str(v: object) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return s


def preview_sheet(ws, *, max_row: int = 120, max_col: int = 30, max_non_empty_rows: int = 60) -> SheetPreview:
    mr = min(int(ws.max_row or 0), max_row)
    mc = min(int(ws.max_column or 0), max_col)

    rows: list[tuple[int, list[str]]] = []
    for r in range(1, mr + 1):
        vals: list[str] = []
        nonempty = 0
        for c in range(1, mc + 1):
            s = _cell_str(ws.cell(r, c).value)
            if s:
                nonempty += 1
            vals.append(s)
        if nonempty:
            rows.append((r, vals))
        if len(rows) >= max_non_empty_rows:
            break

    return SheetPreview(name=str(ws.title), max_row=mr, max_col=mc, non_empty_rows=rows)


def _compact_row(vals: Iterable[str], *, limit: int = 260) -> str:
    parts = [v for v in vals if v]
    s = " | ".join(parts)
    if len(s) > limit:
        s = s[: limit - 3] + "..."
    return s


def main() -> int:
    xlsx = Path("docs") / "Thiết kế đầu báo khói không dây AFE Renesas.xlsx"
    if not xlsx.exists():
        raise SystemExit(f"Missing file: {xlsx}")

    wb = openpyxl.load_workbook(xlsx, data_only=True)

    previews: list[SheetPreview] = []
    for name in wb.sheetnames:
        previews.append(preview_sheet(wb[name]))

    out_md = Path("docs") / "afe_renesas_extract_preview.md"
    with out_md.open("w", encoding="utf-8") as f:
        f.write("# AFE Renesas XLSX — Extract Preview\n\n")
        f.write(f"Source: `{xlsx.as_posix()}`\n\n")
        f.write("## Sheets\n\n")
        for p in previews:
            f.write(f"- {p.name} (scan {p.max_row} rows × {p.max_col} cols)\n")
        f.write("\n")

        for p in previews:
            f.write(f"## {p.name}\n\n")
            if not p.non_empty_rows:
                f.write("(no non-empty rows in preview window)\n\n")
                continue
            for r, vals in p.non_empty_rows:
                compact = _compact_row(vals)
                if compact:
                    f.write(f"- R{r}: {compact}\n")
            f.write("\n")

    print(f"Wrote: {out_md}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
